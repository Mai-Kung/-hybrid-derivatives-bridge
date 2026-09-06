"""Build the Excel trade-log workbook (entry/exit times) for the picked
strategy from results/trade_log.csv.

Sheets:
  - Summary: strategy definition, assumptions, and headline formulas.
  - Period Summary: one row per rebalance (entry/exit time, period return,
    compounding equity curve) -- all via formulas referencing Trade Log.
  - Trade Log: one row per fill (every short position taken).
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path(__file__).parent / "results"
OUT_PATH = RESULTS_DIR / "trade_log.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True)
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="6B7280")
BASE_FONT = Font(name=FONT_NAME)

PCT_FMT = "0.00%"
PRICE_FMT = "General"
DT_FMT = "yyyy-mm-dd hh:mm"


def style_header_row(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def autosize(ws, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build() -> None:
    trades = pd.read_csv(RESULTS_DIR / "trade_log.csv", parse_dates=["entry_time", "exit_time"])
    n_trades = len(trades)
    periods = trades.groupby("entry_time", sort=True).agg(
        exit_time=("exit_time", "first"),
        n_trades=("symbol", "count"),
    ).reset_index()
    n_periods = len(periods)
    trades_per_period = n_trades // n_periods

    wb = Workbook()

    # ---------------- Trade Log sheet ----------------
    ws_t = wb.active
    ws_t.title = "Trade Log"
    headers = ["Entry Time", "Exit Time", "Symbol", "Side", "Signal (chg_30m %)",
               "Entry Price", "Exit Price", "Gross Return %", "Cost %", "Net Return %"]
    ws_t.append(headers)
    style_header_row(ws_t, 1, len(headers))
    for _, r in trades.iterrows():
        cost_pct = r["gross_ret"] - r["net_ret"]
        ws_t.append([
            r["entry_time"].to_pydatetime(), r["exit_time"].to_pydatetime(), r["symbol"], r["side"],
            r["signal_value"] / 100.0, r["entry_price"], r["exit_price"],
            r["gross_ret"], cost_pct, r["net_ret"],
        ])
    for row in range(2, n_trades + 2):
        ws_t.cell(row=row, column=1).number_format = DT_FMT
        ws_t.cell(row=row, column=2).number_format = DT_FMT
        ws_t.cell(row=row, column=5).number_format = PCT_FMT
        ws_t.cell(row=row, column=6).number_format = PRICE_FMT
        ws_t.cell(row=row, column=7).number_format = PRICE_FMT
        ws_t.cell(row=row, column=8).number_format = PCT_FMT
        ws_t.cell(row=row, column=9).number_format = PCT_FMT
        ws_t.cell(row=row, column=10).number_format = PCT_FMT
        for col in range(1, len(headers) + 1):
            ws_t.cell(row=row, column=col).font = BASE_FONT
    autosize(ws_t, {"A": 18, "B": 18, "C": 16, "D": 8, "E": 16, "F": 14, "G": 14, "H": 14, "I": 10, "J": 14})
    ws_t.freeze_panes = "A2"
    last_trade_row = n_trades + 1

    # ---------------- Period Summary sheet ----------------
    ws_p = wb.create_sheet("Period Summary")
    p_headers = ["Period #", "Entry Time", "Exit Time", "# Shorts", "Period Return %",
                 "Equity (start=1.0)", "Running Max", "Drawdown %"]
    ws_p.append(p_headers)
    style_header_row(ws_p, 1, len(p_headers))

    for i, (_, prow) in enumerate(periods.iterrows(), start=1):
        block_start = 2 + (i - 1) * trades_per_period
        block_end = block_start + trades_per_period - 1
        r = i + 1
        ws_p.append([
            i, prow["entry_time"].to_pydatetime(), prow["exit_time"].to_pydatetime(), prow["n_trades"],
            f"=AVERAGE('Trade Log'!J{block_start}:J{block_end})",
            f"=IF(A{r}=1,1+E{r},F{r-1}*(1+E{r}))",
            f"=MAX(G{r-1},F{r})" if i > 1 else f"=F{r}",
            f"=F{r}/G{r}-1",
        ])
    for row in range(2, n_periods + 2):
        ws_p.cell(row=row, column=2).number_format = DT_FMT
        ws_p.cell(row=row, column=3).number_format = DT_FMT
        ws_p.cell(row=row, column=5).number_format = PCT_FMT
        ws_p.cell(row=row, column=6).number_format = "0.0000"
        ws_p.cell(row=row, column=7).number_format = "0.0000"
        ws_p.cell(row=row, column=8).number_format = PCT_FMT
        for col in range(1, len(p_headers) + 1):
            ws_p.cell(row=row, column=col).font = BASE_FONT
    autosize(ws_p, {"A": 10, "B": 18, "C": 18, "D": 10, "E": 15, "F": 16, "G": 14, "H": 14})
    ws_p.freeze_panes = "A2"
    last_period_row = n_periods + 1

    chart = LineChart()
    chart.title = "Equity curve (start = 1.0)"
    chart.y_axis.title = "Equity"
    chart.x_axis.title = "Period #"
    chart.height, chart.width = 10, 22
    data = Reference(ws_p, min_col=6, min_row=1, max_row=last_period_row)
    cats = Reference(ws_p, min_col=1, min_row=2, max_row=last_period_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_p.add_chart(chart, f"J2")

    # ---------------- Summary sheet ----------------
    ws_s = wb.create_sheet("Summary", 0)
    ws_s["A1"] = "Altcoin Perpetual Short Strategy — Trade Log Summary"
    ws_s["A1"].font = TITLE_FONT
    ws_s.merge_cells("A1:D1")

    rows = [
        ("Strategy", "Short the 10 biggest 30-minute gainers each scan (fade-the-pump, short-only)"),
        ("Signal", "chg_30m — Price change %, 30 minutes"),
        ("Direction / Side", "reversion / short_only"),
        ("Basket size (N)", trades_per_period),
        ("Cost assumption", "5 bps per leg (taker + slippage), applied once per trade"),
        ("Universe", "Binance USDT-margined alt perpetuals (excl. BTC/ETH), 24h volume >= $1,000,000"),
        ("Data window", None),
        ("Max snapshot gap allowed", "8 hours (longer gaps -- delistings/data holes -- are dropped)"),
        (None, None),
        ("Total trades (fills)", None),
        ("Total periods (rebalances)", None),
        ("Avg shorts per period", None),
        ("Period return win rate", None),
        ("Best period return", None),
        ("Worst period return", None),
        ("Total return (compounded)", None),
        ("Max drawdown", None),
        ("Periods per year (annualization)", None),
        ("Annualized Sharpe (period returns)", None),
    ]
    start_row = 3
    for i, (label, val) in enumerate(rows):
        r = start_row + i
        if label is not None:
            ws_s.cell(row=r, column=1, value=label).font = LABEL_FONT
        if isinstance(val, str) or (val is not None and not isinstance(val, str)):
            ws_s.cell(row=r, column=2, value=val)

    def r_of(label: str) -> int:
        for i, (lbl, _) in enumerate(rows):
            if lbl == label:
                return start_row + i
        raise KeyError(label)

    ws_s.cell(row=r_of("Data window"), column=2,
              value=(f"=TEXT('Period Summary'!B2,\"yyyy-mm-dd\")&\" to \"&"
                     f"TEXT('Period Summary'!C{last_period_row},\"yyyy-mm-dd\")"))
    ws_s.cell(row=r_of("Total trades (fills)"), column=2, value=f"=COUNTA('Trade Log'!C2:C{last_trade_row})")
    ws_s.cell(row=r_of("Total periods (rebalances)"), column=2, value=f"=COUNTA('Period Summary'!A2:A{last_period_row})")
    ws_s.cell(row=r_of("Avg shorts per period"), column=2, value=f"=AVERAGE('Period Summary'!D2:D{last_period_row})")
    win_row = r_of("Period return win rate")
    ws_s.cell(row=win_row, column=2,
              value=f"=COUNTIF('Period Summary'!E2:E{last_period_row},\">0\")/COUNTA('Period Summary'!E2:E{last_period_row})")
    ws_s.cell(row=win_row, column=2).number_format = PCT_FMT
    best_row = r_of("Best period return")
    ws_s.cell(row=best_row, column=2, value=f"=MAX('Period Summary'!E2:E{last_period_row})")
    ws_s.cell(row=best_row, column=2).number_format = PCT_FMT
    worst_row = r_of("Worst period return")
    ws_s.cell(row=worst_row, column=2, value=f"=MIN('Period Summary'!E2:E{last_period_row})")
    ws_s.cell(row=worst_row, column=2).number_format = PCT_FMT
    total_ret_row = r_of("Total return (compounded)")
    ws_s.cell(row=total_ret_row, column=2, value=f"='Period Summary'!F{last_period_row}-1")
    ws_s.cell(row=total_ret_row, column=2).number_format = PCT_FMT
    dd_row = r_of("Max drawdown")
    ws_s.cell(row=dd_row, column=2, value=f"=MIN('Period Summary'!H2:H{last_period_row})")
    ws_s.cell(row=dd_row, column=2).number_format = PCT_FMT
    ppy_row = r_of("Periods per year (annualization)")
    ws_s.cell(row=ppy_row, column=2,
              value=(f"=COUNTA('Period Summary'!A2:A{last_period_row})/"
                     f"(('Period Summary'!C{last_period_row}-'Period Summary'!B2))*365.25"))
    ws_s.cell(row=ppy_row, column=2).number_format = "0"
    sharpe_row = r_of("Annualized Sharpe (period returns)")
    ws_s.cell(row=sharpe_row, column=2,
              value=(f"=AVERAGE('Period Summary'!E2:E{last_period_row})/"
                     f"STDEV('Period Summary'!E2:E{last_period_row})*SQRT(B{ppy_row})"))
    ws_s.cell(row=sharpe_row, column=2).number_format = "0.00"

    ws_s.cell(row=r_of("Basket size (N)"), column=2).font = BASE_FONT
    for i, (label, _) in enumerate(rows):
        r = start_row + i
        ws_s.cell(row=r, column=2).font = BASE_FONT

    note_row = start_row + len(rows) + 2
    ws_s.cell(row=note_row, column=1,
              value=("Note: this replays historical BI screener snapshots with a fixed cost "
                     "assumption; it is a backtest, not live execution. See ../RESULTS.md for "
                     "full walk-forward validation, cost/liquidity sensitivity, and caveats "
                     "(funding rate not modeled, short-squeeze risk, thin-liquidity slippage)."))
    ws_s.cell(row=note_row, column=1).font = NOTE_FONT
    ws_s.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws_s.row_dimensions[note_row].height = 30
    ws_s.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    autosize(ws_s, {"A": 34, "B": 46})

    wb.save(OUT_PATH)
    print(f"Saved {OUT_PATH} ({n_trades} trades, {n_periods} periods)")


if __name__ == "__main__":
    build()
